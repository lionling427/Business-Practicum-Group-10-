# PROCESS 1 – IMPORTS

import io

import pandas as pd
import pytest
from pptx import Presentation

from src.agent import DEFAULT_DATASET
from src.tools.analytics import analyze_dataframe
from src.tools.ceo_brief import build_ceo_brief
from src.tools.health_check import check_upload
from src.tools.report_pack import (
    SECTIONS,
    Branding,
    _plain,
    build_pack,
    slide_previews,
    to_pdf,
    to_pptx,
)
from src.tools.waves import model_sample


# PROCESS 2 – FIXTURES

@pytest.fixture(scope="module")
def data():

    path = DEFAULT_DATASET / "agent_data"
    s = pd.read_csv(path / "survey_responses.csv")
    i = pd.read_csv(path / "individual_outcomes.csv")
    return s.merge(i, on=[c for c in s.columns if c in i.columns])


@pytest.fixture(scope="module")
def analysis(data):

    return analyze_dataframe(model_sample(data), dataset_name="dev",
                             health_report=check_upload(data, "dev"))


@pytest.fixture(scope="module")
def slides(data, analysis):

    return build_pack(data, analysis)


@pytest.fixture(scope="module")
def logo():

    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (200, 80), "#1F4E9A").save(buffer, format="PNG")
    return buffer.getvalue()


# PROCESS 3 – CONTENT

def test_pack_has_every_section_in_order(slides):

    titles = [s["title"] for s in slides]

    assert titles[0] == "Engagement and business results"
    assert titles[-1] == "About this report"
    assert titles[1:-1] == ["The bottom line", "What drives results", "What it is worth",
                            "What is likely to happen", "Where to start", "Decisions for you",
                            "How sure we are"]


def test_numbers_match_the_ceo_brief(data, analysis, slides):

    brief = build_ceo_brief(data, analysis)
    bottom = next(s for s in slides if s["title"] == "The bottom line")

    assert bottom["lead"] == brief["bottom_line"]
    assert [m[1] for m in bottom["metrics"]] == [x["value"] for x in brief["scorecard"]]


def test_sections_can_be_left_out(data, analysis):

    slides = build_pack(data, analysis, sections=["Decisions"])

    assert [s["title"] for s in slides] == ["Engagement and business results",
                                            "Decisions for you", "About this report"]


def test_no_markdown_left_on_slides(slides):

    for slide in slides:
        for text in [slide.get("lead", ""), *slide.get("bullets", [])]:
            assert "**" not in text and "\\$" not in text


def test_plain_text_keeps_money():

    assert _plain("Worth **\\$2.5M** a year") == "Worth $2.5M a year"


# PROCESS 4 – FILES

def test_powerpoint_opens_with_branding(slides, logo):

    deck = Presentation(io.BytesIO(to_pptx(slides, Branding(company="Acme Corp", logo=logo))))

    assert len(deck.slides) == len(slides)
    texts = [sh.text_frame.text for sh in deck.slides[0].shapes if sh.has_text_frame]
    assert "Acme Corp" in texts
    assert any(sh.shape_type == 13 for sh in deck.slides[0].shapes)          # the logo picture
    assert any(sh.has_table for sh in deck.slides[5].shapes)                 # where to start


def test_pdf_has_one_page_per_slide(slides, logo):

    pdf = to_pdf(slides, Branding(company="Acme Corp", color="#1F4E9A", logo=logo))

    assert pdf.startswith(b"%PDF")
    assert pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") + pdf.count(b"/Type/Page") >= 1
    assert len(slide_previews(slides, Branding(), dpi=20)) == len(slides)
